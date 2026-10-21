import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    PRIMARY = RGBColor(18, 59, 93)     # #123B5D - Deep Ocean Blue
    TEAL = RGBColor(19, 184, 166)      # #13B8A6 - Tropical Teal
    ORANGE = RGBColor(245, 158, 66)    # #F59E42 - Sunset Orange
    DARK = RGBColor(23, 43, 77)        # #172B4D - Slate Dark
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC - Off white
    CARD_BG = RGBColor(255, 255, 255)  # White
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    MUTED = RGBColor(100, 116, 139)    # Slate 500
    WHITE = RGBColor(255, 255, 255)
    SUCCESS = RGBColor(22, 163, 74)    # Green

    blank_layout = prs.slide_layouts[6] # Blank slide

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="TRAVELMATE PROJECT OVERVIEW"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = PRIMARY
        banner.line.color.rgb = PRIMARY

        # Accent bar
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.color.rgb = TEAL

        # Title text box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEAL

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

    def add_footer(slide, slide_num):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.4))
        tf = tb.text_frame
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = f"TravelMate – Smart Tourist Day-Visit Planner  |  ITE2953 Software Project  |  Slide {slide_num} of 5"
        p.font.size = Pt(10)
        p.font.color.rgb = MUTED

    # ==========================================
    # SLIDE 0: TITLE SLIDE
    # ==========================================
    s0 = prs.slides.add_slide(blank_layout)
    set_slide_background(s0, PRIMARY)

    # Decorative shapes
    dec1 = s0.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(10.5), Inches(0), Inches(2.833), Inches(7.5))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = RGBColor(15, 45, 75)
    dec1.line.fill.background()

    # Top badge
    badge = s0.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.6), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = TEAL
    badge.line.fill.background()
    btf = badge.text_frame
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "ITE2953 SOFTWARE PROJECT • FINAL DEFENSE"
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = PRIMARY

    # Title & Subtitle Box
    tbox = s0.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(9.5), Inches(2.4))
    ttf = tbox.text_frame
    ttf.word_wrap = True

    tp1 = ttf.paragraphs[0]
    tp1.text = "TravelMate"
    tp1.font.size = Pt(48)
    tp1.font.bold = True
    tp1.font.color.rgb = WHITE

    tp2 = ttf.add_paragraph()
    tp2.text = "Smart Tourist Day-Visit Planner"
    tp2.font.size = Pt(30)
    tp2.font.bold = True
    tp2.font.color.rgb = ORANGE

    tp3 = ttf.add_paragraph()
    tp3.text = "An Enterprise Full-Stack Web Solution for Tailored Day-Trip Exploration and Optimized Itinerary Generation (25km Radius from Atulugama, Sri Lanka)"
    tp3.font.size = Pt(14)
    tp3.font.color.rgb = RGBColor(203, 213, 225)

    # Info Card
    info_card = s0.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.5), Inches(8.5), Inches(2.2))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = RGBColor(24, 75, 115)
    info_card.line.color.rgb = TEAL
    itf = info_card.text_frame
    itf.margin_left = Inches(0.3)
    itf.margin_top = Inches(0.25)

    ip0 = itf.paragraphs[0]
    ip0.text = "PROJECT METADATA & CANDIDATE DETAILS"
    ip0.font.size = Pt(11)
    ip0.font.bold = True
    ip0.font.color.rgb = TEAL

    details = [
        ("Candidate Name", "Rukshana Roshan"),
        ("Registration Number", "E2410132"),
        ("Degree Programme", "B.Sc. (Hons) in Information Technology"),
        ("Academic Institution", "University of Moratuwa, Sri Lanka"),
        ("Course Module", "ITE2953 – Software Project (Weightage: 60% Final Mark)"),
        ("Technology Stack", "Java 17, Spring Boot 3.2, MySQL 8, Bootstrap 5, StarUML")
    ]
    for label, val in details:
        p = itf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"•  {label}: "
        run1.font.bold = True
        run1.font.size = Pt(11)
        run1.font.color.rgb = WHITE
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(11)
        run2.font.color.rgb = RGBColor(226, 232, 240)

    # Speaker Notes Slide 0
    s0.notes_slide.notes_text_frame.text = (
        "Good morning / afternoon, respected examiners and faculty members. "
        "My name is Rukshana Roshan, Registration Number E2410132. Today, I am proud to present my final software engineering project "
        "for course module ITE2953: TravelMate – Smart Tourist Day-Visit Planner. Over the next five slides, I will provide a quick, "
        "comprehensive overview of our problem domain, layered system architecture, nearest-neighbour optimization algorithm, database design, "
        "and automated testing results."
    )

    # ==========================================
    # SLIDE 1: PROBLEM DOMAIN & OBJECTIVES (15 pts)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, LIGHT_BG)
    add_header(s1, "Problem Domain, Project Objectives & Stakeholders", "RUBRIC: PROJECT UNDERSTANDING & PROBLEM DOMAIN")
    add_footer(s1, 1)

    # Left Column: Problem Statement & Stakeholders
    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf1 = c1.text_frame
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.3)

    p = tf1.paragraphs[0]
    p.text = "1. Problem Domain & Background"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    points_p1 = [
        ("Information Fragmentation", "Tourists visiting Atulugama, Bandaragama, and Kalutara face scattered, outdated visiting hours, ticket fees, and travel distance data across disparate social media pages."),
        ("Unrealistic Day Schedules", "Visitors struggle to pace one-day trips. Standard navigation tools calculate point-to-point driving but ignore site dwell times, meal stops, and local road-winding factors."),
        ("Target Stakeholders", "Domestic day tourists, visiting families, adventure sports enthusiasts, and local administrative authorities managing regional tourism assets.")
    ]
    for title, desc in points_p1:
        p_t = tf1.add_paragraph()
        p_t.space_before = Pt(8)
        run = p_t.add_run()
        run.text = f"• {title}: "
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = DARK
        run2 = p_t.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = MUTED

    # Right Column: Core Project Objectives & Scope
    c2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf2 = c2.text_frame
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "2. Specific Objectives & Scope"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    points_p2 = [
        ("Curated 25km Spatial Radius", "Filter and present 10 verified high-interest attractions across 6 distinct categories centered at Atulugama (6.7022° N, 80.0528° E)."),
        ("Algorithmic Timeline Generator", "Generate conflict-free, clock-timed itineraries incorporating transit buffers, stay durations, and automated midday lunch break scheduling."),
        ("Enterprise Administrative Governance", "Full role-gated administrative suite for CRUD operations on destinations, categories, and moderation of tourist inquiries."),
        ("Decoupled Multi-Platform Deployment", "RESTful backend ready for cloud containerization (Docker on Render) and responsive static frontend deployable on CDN / GitHub Pages.")
    ]
    for title, desc in points_p2:
        p_t = tf2.add_paragraph()
        p_t.space_before = Pt(8)
        run = p_t.add_run()
        run.text = f"• {title}: "
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = DARK
        run2 = p_t.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = MUTED

    # Speaker Notes Slide 1
    s1.notes_slide.notes_text_frame.text = (
        "On Slide 1, we highlight the core problem domain and objectives. Travelers visiting the Atulugama and Kalutara region "
        "frequently suffer from fragmented information—incorrect entrance fees, unverified opening hours, and unrealistic day schedules "
        "that don't account for local traffic or stay durations. TravelMate resolves this by anchoring an accurate 25km geographic catalog "
        "around Atulugama, providing a deterministic nearest-neighbour itinerary planner that constructs a realistic, timed schedule, "
        "and offering local tourism administrators a dedicated portal to maintain verified data."
    )

    # ==========================================
    # SLIDE 2: SYSTEM ARCHITECTURE & TECH STACK (25 pts)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "Layered Architecture, Technology Stack & API Contract", "RUBRIC: TECHNICAL IMPLEMENTATION KNOWLEDGE")
    add_footer(s2, 2)

    # 3 Column Cards for the Layered Architecture
    col_w = Inches(3.64)
    gap = Inches(0.38)
    left_start = Inches(0.8)

    # Card 1: Frontend Layer
    c_fe = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start, Inches(1.4), col_w, Inches(5.4))
    c_fe.fill.solid()
    c_fe.fill.fore_color.rgb = CARD_BG
    c_fe.line.color.rgb = CARD_BORDER
    tfe = c_fe.text_frame
    tfe.margin_left = tfe.margin_right = tfe.margin_top = Inches(0.25)
    p = tfe.paragraphs[0]
    p.text = "Presentation Layer"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL

    p_sub = tfe.add_paragraph()
    p_sub.text = "Decoupled Modern Client"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = MUTED

    fe_bullets = [
        ("HTML5 & Bootstrap 5", "Responsive mobile-first grid, accessible navigation, and modal dialogues."),
        ("Vanilla ES6+ JavaScript", "Clean client logic without heavyweight framework bloat."),
        ("Unified API Client (api.js)", "Centralized TM_API object wrapping fetch requests with error handling."),
        ("Interactive Maps Integration", "Direct integration with Google Maps multi-stop routing API."),
        ("Print / PDF Export", "Dedicated CSS @media print styling for offline physical trip itineraries.")
    ]
    for t, d in fe_bullets:
        pt = tfe.add_paragraph()
        pt.space_before = Pt(8)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Card 2: Backend REST Layer
    c_be = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start + col_w + gap, Inches(1.4), col_w, Inches(5.4))
    c_be.fill.solid()
    c_be.fill.fore_color.rgb = CARD_BG
    c_be.line.color.rgb = CARD_BORDER
    tbe = c_be.text_frame
    tbe.margin_left = tbe.margin_right = tbe.margin_top = Inches(0.25)
    p = tbe.paragraphs[0]
    p.text = "Backend REST Layer"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p_sub = tbe.add_paragraph()
    p_sub.text = "Spring Boot 3.2 & Java 17"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = MUTED

    be_bullets = [
        ("Clean Layered Hierarchy", "Controller → Service Interface → ServiceImpl → Repository."),
        ("DTO Encapsulation", "Zero JPA entity leakage across API boundaries; validated with Bean Validation (@Valid)."),
        ("ApiResponse<T> Envelope", "Predictable response structure: success, message, data, timestamp."),
        ("Global Exception Handling", "Centralized @ControllerAdvice mapping domain exceptions to RFC HTTP status codes."),
        ("Security Crypto", "BCrypt salted password hashing with Spring Security Crypto.")
    ]
    for t, d in be_bullets:
        pt = tbe.add_paragraph()
        pt.space_before = Pt(8)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Card 3: Persistence & DevOps Layer
    c_db = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start + (col_w + gap)*2, Inches(1.4), col_w, Inches(5.4))
    c_db.fill.solid()
    c_db.fill.fore_color.rgb = CARD_BG
    c_db.line.color.rgb = CARD_BORDER
    tdb = c_db.text_frame
    tdb.margin_left = tdb.margin_right = tdb.margin_top = Inches(0.25)
    p = tdb.paragraphs[0]
    p.text = "Data & Deployment"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ORANGE

    p_sub = tdb.add_paragraph()
    p_sub.text = "Relational Storage & Cloud"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = MUTED

    db_bullets = [
        ("MySQL 8.0 Relational DB", "7 normalized tables enforcing referential integrity and cascading deletions."),
        ("Spring Data JPA & Hibernate", "High-performance ORM with HikariCP connection pooling."),
        ("Dual Test Strategy", "In-memory H2 database runs 38 automated test cases in CI/CD without live MySQL dependency."),
        ("Docker Containerization", "Optimized multi-stage Dockerfile for fast, minimal container builds."),
        ("Cloud Ready (Render / Pages)", "Decoupled hosting with environment auto-detection via config.js.")
    ]
    for t, d in db_bullets:
        pt = tdb.add_paragraph()
        pt.space_before = Pt(8)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Speaker Notes Slide 2
    s2.notes_slide.notes_text_frame.text = (
        "Slide 2 showcases our technical implementation and architecture. The solution implements a clean four-tier layered architecture. "
        "Controllers are lightweight and map HTTP requests to services, wrapping responses in an ApiResponse envelope. DTOs completely isolate "
        "database entities so database schemas are never exposed across the network. Security is enforced through BCrypt password hashing. "
        "The persistence layer utilizes MySQL 8 in production and an in-memory H2 database for our automated test suite."
    )

    # ==========================================
    # SLIDE 3: FUNCTIONALITY & NEAREST-NEIGHBOUR ALGORITHM (25 pts)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "Core System Functionality & Nearest-Neighbour Routing Algorithm", "RUBRIC: SYSTEM DEMONSTRATION & FUNCTIONALITY")
    add_footer(s3, 3)

    # Left: Core Features (Tourist & Admin)
    cf_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    cf_box.fill.solid()
    cf_box.fill.fore_color.rgb = CARD_BG
    cf_box.line.color.rgb = CARD_BORDER
    tcf = cf_box.text_frame
    tcf.margin_left = tcf.margin_right = tcf.margin_top = Inches(0.3)

    p = tcf.paragraphs[0]
    p.text = "Implemented System Capabilities"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    feats = [
        ("Catalog Discovery", "Multi-criteria filtering by category chips, real-time keyword search, and sorting by duration, distance, and rating."),
        ("Rich Destination Profiles", "High-resolution photos, verified GPS coordinates, visiting duration, admission fee, opening hours, and travel etiquette tips."),
        ("Trip Shortlisting", "One-click 'Add to Day Plan' with duplicate prevention toast notifications and live badge counter."),
        ("Admin Portal & CRUD", "Secured dashboard with KPI summary cards, modal forms for adding/editing destinations, and category moderation."),
        ("Public Inquiries", "Tourists submit inquiries via contact form; admins review, filter, and mark messages as read.")
    ]
    for t, d in feats:
        pt = tcf.add_paragraph()
        pt.space_before = Pt(6)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Right: Algorithm Steps Box
    algo_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    algo_box.fill.solid()
    algo_box.fill.fore_color.rgb = CARD_BG
    algo_box.line.color.rgb = CARD_BORDER
    talgo = algo_box.text_frame
    talgo.margin_left = talgo.margin_right = talgo.margin_top = Inches(0.3)

    p = talgo.paragraphs[0]
    p.text = "Smart Nearest-Neighbour Itinerary Algorithm"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL

    steps = [
        ("Step 1: Anchor Origin", "Departure is always anchored at Atulugama Central Coordinates (6.7022° N, 80.0528° E)."),
        ("Step 2: Road Distance Matrix", "Computes Haversine great-circle distance multiplied by a 1.25× empirical road-winding coefficient for Sri Lankan suburban terrain."),
        ("Step 3: Realistic Transit Speeds", "Applies vehicle travel speeds (Car: 35 km/h, Motorcycle: 40 km/h) to compute minute-accurate inter-stop transit buffers."),
        ("Step 4: Greedy Nearest-Neighbour", "Iteratively selects the closest unvisited shortlisted stop, minimizing cumulative transit fatigue."),
        ("Step 5: Time Slots & Lunch Break", "Calculates cumulative arrival and departure clock times. If the tour spans 12:30 PM – 1:30 PM, a 60-minute lunch break is inserted."),
        ("Step 6: Overflow Boundary Check", "If a destination cannot be completed before the user-specified Finish Time, it is gracefully flagged as an advisory overflow stop.")
    ]
    for t, d in steps:
        pt = talgo.add_paragraph()
        pt.space_before = Pt(5)
        r1 = pt.add_run()
        r1.text = f"{t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Speaker Notes Slide 3
    s3.notes_slide.notes_text_frame.text = (
        "Slide 3 details our system functionality and itinerary algorithm. Rather than relying on an opaque, non-deterministic black box, "
        "we designed a deterministic nearest-neighbour routing algorithm tailored to local conditions. The algorithm anchors the trip at Atulugama, "
        "applies a 1.25 road-winding factor to the Haversine distance, applies realistic vehicle speeds (35 km/h for cars), builds a minute-by-minute "
        "chronological timeline, inserts a lunch break between 12:30 and 1:30 PM, and flags any stops that exceed the requested finish time."
    )

    # ==========================================
    # SLIDE 4: DATABASE & DESIGN DECISIONS (15 pts)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Relational Database Design & Key Software Engineering Decisions", "RUBRIC: SOFTWARE ENGINEERING & DESIGN DECISIONS")
    add_footer(s4, 4)

    # Left: Database Entities & Schema
    db_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    db_card.fill.solid()
    db_card.fill.fore_color.rgb = CARD_BG
    db_card.line.color.rgb = CARD_BORDER
    tdbc = db_card.text_frame
    tdbc.margin_left = tdbc.margin_right = tdbc.margin_top = Inches(0.3)

    p = tdbc.paragraphs[0]
    p.text = "Relational Schema (7 Tables)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    tables = [
        ("users", "User authentication, hashed passwords, roles (ADMIN / TOURIST)."),
        ("categories", "Attraction taxonomy (Recreation, Adventure, Nature, Beaches, etc.)."),
        ("attractions", "10 seeded destinations with GPS coords, durations, opening hours, fees."),
        ("itinerary", "Master day plan record linked to user, date, start/end time, travel mode."),
        ("itinerary_items", "Ordered stops per plan with arrival, departure, transit buffer, distance."),
        ("favorites", "Many-to-Many junction table bookmarking saved attractions per user."),
        ("contact_messages", "Tourist inquiry messages with read/unread moderation status.")
    ]
    for t, d in tables:
        pt = tdbc.add_paragraph()
        pt.space_before = Pt(5)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Right: Justified Engineering Decisions
    eng_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    eng_card.fill.solid()
    eng_card.fill.fore_color.rgb = CARD_BG
    eng_card.line.color.rgb = CARD_BORDER
    teng = eng_card.text_frame
    teng.margin_left = teng.margin_right = teng.margin_top = Inches(0.3)

    p = teng.paragraphs[0]
    p.text = "Justified Design & Engineering Decisions"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL

    decisions = [
        ("Deterministic Nearest-Neighbour", "Guarantees mathematically explainable, reproducible itineraries without unpredictable hallucinations or third-party paid API quotas."),
        ("Standard Java POJOs (No Lombok)", "Prevents Lombok annotation-processor build incompatibilities across varying JDK versions (Java 17, 21, and 26), ensuring 100% build reproducibility for examiners."),
        ("Dual-Layer Data Resilience", "The frontend features a client-side offline dataset (data.js) coupled with live REST synchronization, ensuring the system runs smoothly even in offline grading environments."),
        ("Strict DTO Encapsulation", "Prevents JPA serialization recursion and over-posting security vulnerabilities by keeping entity models strictly within the service layer."),
        ("Database Cascades", "Foreign key constraints utilize ON DELETE CASCADE where appropriate, preventing orphaned child records.")
    ]
    for t, d in decisions:
        pt = teng.add_paragraph()
        pt.space_before = Pt(6)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Speaker Notes Slide 4
    s4.notes_slide.notes_text_frame.text = (
        "On Slide 4, we examine our software engineering and architectural choices. Our database schema consists of seven normalized tables "
        "designed in 3rd Normal Form with cascading foreign key integrity. Key engineering choices include using a deterministic algorithm "
        "for predictable tour plans, choosing standard Java POJOs over Lombok to ensure seamless cross-platform compilation on examiners' machines, "
        "implementing strict DTO encapsulation to avoid security over-posting, and dual-layer data resilience so the application remains robust."
    )

    # ==========================================
    # SLIDE 5: TESTING & VERIFICATION RESULTS (10 pts)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Software Testing Strategy, Automated Results & Project Deliverables", "RUBRIC: COMMUNICATION & TESTING RIGOR")
    add_footer(s5, 5)

    # Top Metric Banner (3 Stat Boxes)
    box_w = Inches(3.64)
    stat1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), box_w, Inches(1.2))
    stat1.fill.solid()
    stat1.fill.fore_color.rgb = CARD_BG
    stat1.line.color.rgb = CARD_BORDER
    st1 = stat1.text_frame
    st1.margin_top = Inches(0.15)
    st1.paragraphs[0].text = "38 / 38 TESTS PASSED"
    st1.paragraphs[0].font.size = Pt(18)
    st1.paragraphs[0].font.bold = True
    st1.paragraphs[0].font.color.rgb = SUCCESS
    st1.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = st1.add_paragraph()
    p.text = "100% Automated Maven Test Pass Rate"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER

    stat2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + box_w + gap, Inches(1.4), box_w, Inches(1.2))
    stat2.fill.solid()
    stat2.fill.fore_color.rgb = CARD_BG
    stat2.line.color.rgb = CARD_BORDER
    st2 = stat2.text_frame
    st2.margin_top = Inches(0.15)
    st2.paragraphs[0].text = "20 / 20 TEST CASES"
    st2.paragraphs[0].font.size = Pt(18)
    st2.paragraphs[0].font.bold = True
    st2.paragraphs[0].font.color.rgb = PRIMARY
    st2.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = st2.add_paragraph()
    p.text = "University Specification (TC01–TC20)"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER

    stat3 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (box_w + gap)*2, Inches(1.4), box_w, Inches(1.2))
    stat3.fill.solid()
    stat3.fill.fore_color.rgb = CARD_BG
    stat3.line.color.rgb = CARD_BORDER
    st3 = stat3.text_frame
    st3.margin_top = Inches(0.15)
    st3.paragraphs[0].text = "0 FAILURES / 0 REGRESSIONS"
    st3.paragraphs[0].font.size = Pt(18)
    st3.paragraphs[0].font.bold = True
    st3.paragraphs[0].font.color.rgb = TEAL
    st3.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = st3.add_paragraph()
    p.text = "Spring Boot MockMvc & JUnit 5 Suite"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER

    # Bottom Left: Testing Layers Breakdown
    btm_left = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.8), Inches(5.6), Inches(4.0))
    btm_left.fill.solid()
    btm_left.fill.fore_color.rgb = CARD_BG
    btm_left.line.color.rgb = CARD_BORDER
    tbl = btm_left.text_frame
    tbl.margin_left = tbl.margin_right = tbl.margin_top = Inches(0.25)

    p = tbl.paragraphs[0]
    p.text = "Testing Strategies Implemented"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    t_strat = [
        ("Unit Testing (Mockito)", "18 tests isolating service logic, DTO mapping, and edge condition handling."),
        ("Integration Testing (MockMvc + H2)", "20 tests verifying REST routes, JSON payloads, and validation constraints against in-memory H2 database."),
        ("Security Regression Testing", "Verified BCrypt login authentication and 401 Unauthorized handling for unauthenticated requests."),
        ("Cross-Device UI Verification", "Verified responsive rendering across 375px mobile, tablet, and desktop breakpoints.")
    ]
    for t, d in t_strat:
        pt = tbl.add_paragraph()
        pt.space_before = Pt(5)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Bottom Right: Submission Deliverables Summary
    btm_right = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.8), Inches(5.7), Inches(4.0))
    btm_right.fill.solid()
    btm_right.fill.fore_color.rgb = CARD_BG
    btm_right.line.color.rgb = CARD_BORDER
    tbr = btm_right.text_frame
    tbr.margin_left = tbr.margin_right = tbr.margin_top = Inches(0.25)

    p = tbr.paragraphs[0]
    p.text = "Packaged Submission Deliverables"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEAL

    delivs = [
        ("Developed System Source Code", "Complete Maven backend (Spring Boot 3.2) + frontend client with seed data."),
        ("StarUML Design Project (.mdj)", "All 5 required design diagrams: Use Case, Activity, Sequence, Class, and ER/EER."),
        ("5-Slide Presentation Deck", "Supplied in native PPTX, interactive HTML, and Markdown format."),
        ("Video Demonstration Script", "14-stage timed walkthrough script covering all major features."),
        ("Test Verification Report", "Detailed TC01–TC20 traceability matrix for Final Report Chapter 5."),
        ("Single Zipped Submission Folder", "Packaged and ready for upload to the course submission portal.")
    ]
    for t, d in delivs:
        pt = tbr.add_paragraph()
        pt.space_before = Pt(5)
        r1 = pt.add_run()
        r1.text = f"• {t}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK
        r2 = pt.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED

    # Speaker Notes Slide 5
    s5.notes_slide.notes_text_frame.text = (
        "In conclusion, on Slide 5, we present our testing verification and deliverables. All 38 automated test cases in our test suite "
        "execute with zero errors and zero failures, covering all 20 university test cases. The project is fully realized and packaged into "
        "our submission folder, complete with the developed system, StarUML design models, presentation slides, demonstration script, "
        "and test report. Thank you for your time, and I welcome any questions regarding the implementation."
    )

    # Save presentation
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out = os.path.abspath("docs/TravelMate_Presentation.pptx")
    build_presentation(out)
